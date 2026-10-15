"""Category axis must not have a numeric zero minimum / extra blank category."""
from pathlib import Path
import tempfile
from openpyxl import load_workbook
from tensile_core import process_files

paths=sorted((Path(__file__).parent/'examples/synthetic_data').glob('*.csv'))
with tempfile.TemporaryDirectory() as folder:
    output=Path(folder)/'categories.xlsx'
    process_files([(p,4,.01,20,{}) for p in paths],output,emit=lambda _:None)
    wb=load_workbook(output);chart=wb['Properties']._charts[0]
    for c in chart._charts:
        assert c.x_axis.scaling.min is None
        assert c.x_axis.scaling.max is None
        assert c.y_axis.scaling.min==0
        for series in c.series:
            assert series.cat.numRef.f==f"'Properties'!$W$2:$W${len(paths)+1}"
    assert [wb['Properties'].cell(i,23).value for i in range(2,len(paths)+2)]==[p.stem for p in paths]
print('Category range matches samples; no zero-index category; stress/modulus start at zero')
