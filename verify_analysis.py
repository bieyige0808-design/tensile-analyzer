"""Numerical checks of modulus, fracture and smoothing invariance."""
import tempfile
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from tensile_core import analyze_trial,process_files
from examples.generate_synthetic_data import generate

t=np.arange(0,2.01,.01);strain=t*.01;stress=5000*strain+2
stress[strain>.015]=0
r=analyze_trial(np.column_stack((t,stress*.04,strain*20)),4,.01,20)
assert abs(r['modulus']-5)<1e-8
assert abs(r['break_strain']-1.5)<1e-8,(r['break_strain'],r['break_index'])
assert abs(r['toughness']-(2500*.015**2+2*.015))<1e-8
stress=5000*strain+2
no_break=analyze_trial(np.column_stack((t,stress*.04,strain*20)),4,.01,20)
assert no_break['break_strain'] is None and no_break['toughness'] is None
with tempfile.TemporaryDirectory() as folder:
    generate(folder)
    configs=[(p,4,.01,20,{}) for p in sorted(Path(folder).glob('*.csv'))]
    a=process_files(configs,Path(folder)/'raw.xlsx',smooth=False,emit=lambda _:None)
    b=process_files(configs,Path(folder)/'smooth.xlsx',smooth=True,emit=lambda _:None)
    for sample in a:
        for (_,x),(_,y) in zip(a[sample],b[sample]):
            for k in ('strength','modulus','break_strain','toughness'):assert x[k]==y[k]
    wb=load_workbook(Path(folder)/'smooth.xlsx')
    assert wb.sheetnames[:3]==['Properties','Stress_Strain','Summary']
    assert len(wb['Properties']._charts)==1 and len(wb['Stress_Strain']._charts)==1
    assert len(wb['Summary']['A'])==11
print('All numerical and workbook checks passed')
