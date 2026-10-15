"""Tensile analysis of instrument CSV files; values use engineering strain."""
from pathlib import Path
import csv
import numpy as np
from scipy.signal import savgol_filter
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, BarChart, Reference, Series
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.styles import Font, PatternFill
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import LegendEntry
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun

def read_trials(path):
    for encoding in ('utf-8-sig','cp932'):
        try:
            with Path(path).open(encoding=encoding,newline='') as f:
                rows=list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError('Unsupported CSV encoding')
    if len(rows)<4 or len(rows[0])%3:
        raise ValueError('Expected three header rows and time/force/stroke column groups')
    trials=[]
    for col in range(0,len(rows[0]),3):
        values=[]
        for row in rows[3:]:
            try:
                triple=[float(row[col+k]) for k in range(3)]
                if np.isfinite(triple).all(): values.append(triple)
            except (ValueError,IndexError): continue
        if values:
            trials.append((rows[0][col] or str(col//3+1),np.array(values)))
    if not trials: raise ValueError('No numeric trials found')
    return trials

def analyze_trial(raw,width,thickness,length,fit_min=.05,fit_max=.25,drop=.5,override=None):
    if min(width,thickness,length)<=0: raise ValueError('Dimensions must be positive')
    if not 0<=fit_min<fit_max or not 0<drop<1: raise ValueError('Invalid fitting range or drop fraction')
    time,force,stroke=raw.T
    if np.any(np.diff(time)<0) or np.any(np.diff(stroke)<-1e-6):
        raise ValueError('Time/stroke must increase; check the selected CSV')
    stress=force/(width*thickness)
    strain=stroke/length
    peak=float(stress.max())
    break_idx=None
    if override is not None:
        if not np.isfinite(override) or override<strain.min()*100 or override>strain.max()*100:
            raise ValueError('Manual fracture strain must be within the measured strain range')
        break_idx=int(np.argmin(abs(strain*100-override)))
        status='manual'
    else:
        # Search for the start of a sharp drop and require a persistent low load.
        for i in range(len(time)-3):
            if stress[i]<peak*.2: continue
            stop=min(len(time),int(np.searchsorted(time,time[i]+.1,side='right')))
            hits=np.flatnonzero(stress[i+1:stop]<=stress[i]*(1-drop))
            if not len(hits): continue
            j=i+1+int(hits[0])
            end=min(len(time),j+5)
            if end-j<3 or np.any(stress[j:end]>stress[i]*(1-drop)): continue
            # Locate the final local high point before the confirmed decline.
            break_idx=i+int(np.argmax(stress[i:j]))
            break
        status='detected' if break_idx is not None else 'not detected'
    upto=break_idx+1 if break_idx is not None else len(stress)
    mask=(strain*100>=fit_min)&(strain*100<=fit_max)&(np.arange(len(strain))<upto)
    if mask.sum()<3 or np.ptp(strain[mask])<=0:
        raise ValueError('Fewer than three points in the modulus fitting range')
    slope,intercept=np.polyfit(strain[mask],stress[mask],1)
    predicted=slope*strain[mask]+intercept
    sst=float(np.sum((stress[mask]-stress[mask].mean())**2))
    r2=1-float(np.sum((stress[mask]-predicted)**2))/sst if sst else None
    return dict(raw=raw,strain=strain,stress=stress,break_index=break_idx,
        break_status=status,strength=float(stress[:upto].max()),modulus=float(slope/1000),
        break_strain=float(strain[break_idx]*100) if break_idx is not None else None,
        toughness=float(np.trapezoid(stress[:upto],strain[:upto])) if break_idx is not None else None,
        r2=r2,intercept=float(intercept),fit_min=fit_min,fit_max=fit_max,
        width=width,thickness=thickness,length=length)

def text_style():
    cp=CharacterProperties(latin=DFont(typeface='Arial'),sz=1400,b=False)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),endParaRPr=cp)])

def format_chart(chart,x_title,y_title):
    chart.width=12.7; chart.height=12.7; chart.title=None
    chart.x_axis.title=x_title; chart.y_axis.title=y_title
    chart.layout=Layout(manualLayout=ManualLayout(layoutTarget='inner',xMode='edge',yMode='edge',wMode='factor',hMode='factor',x=.16,y=.04,w=.80,h=.80))
    chart.graphical_properties=GraphicalProperties(solidFill='FFFFFF')
    chart.graphical_properties.ln.noFill=True
    chart.plot_area.spPr=GraphicalProperties(solidFill='FFFFFF')
    chart.plot_area.spPr.ln.solidFill='000000'; chart.plot_area.spPr.ln.w=12700
    for axis in (chart.x_axis,chart.y_axis):
        axis.txPr=text_style(); axis.majorGridlines=None; axis.minorGridlines=None
        axis.spPr=GraphicalProperties(); axis.spPr.ln.solidFill='000000'; axis.spPr.ln.w=12700
        axis.scaling.min=0; axis.majorTickMark='in'; axis.minorTickMark='none'; axis.crosses='min'
        axis.delete=False;axis.tickLblPos='nextTo';axis.numFmt='0.##'
        axis.crossesAt=None
        if axis.title:
            axis.title.txPr=text_style()
            for p in axis.title.tx.rich.p:
                p.pPr=ParagraphProperties(defRPr=CharacterProperties(latin=DFont(typeface='Arial'),sz=1400,b=False))
                for run in p.r:
                    run.rPr=CharacterProperties(latin=DFont(typeface='Arial'),sz=1400,b=False)
    chart.x_axis.axPos='b';chart.y_axis.axPos='l'
    if chart.legend: chart.legend.txPr=text_style(); chart.legend.position='b';chart.legend.overlay=False

def included_trials(trials):
    return [(label,r) for label,r in trials if r.get('included',True)]

def sample_statistics(trials):
    eligible=included_trials(trials)
    stats={}
    for key in ('strength','modulus','break_strain','toughness'):
        nums=[r[key] for _,r in eligible if r[key] is not None]
        stats[key]=(float(np.mean(nums)) if nums else None,float(np.std(nums,ddof=1)) if len(nums)>1 else None)
    return stats

def plot_values(result,smooth):
    y=result['stress'].copy()
    end=result['break_index']+1 if result['break_index'] is not None else len(y)
    # Do not let smoothing spread the post-fracture drop into the pre-fracture curve.
    if smooth and end>=11: y[:end]=savgol_filter(y[:end],11,3)
    return y,end

def export_excel(samples,path,smooth=False):
    wb=Workbook(); bars=wb.active; bars.title='Properties'
    overview=wb.create_sheet('Stress_Strain'); summary=wb.create_sheet('Summary')
    headers=['Sample','Trial','Break status','Strength (MPa)','Modulus (GPa)','Break strain (%)','Toughness (MJ/m3)','Fit R2','Width (mm)','Thickness (mm)','Gauge length (mm)','Fit min (%)','Fit max (%)','Included in mean']
    summary.append(headers)
    barheaders=['Sample','Strength mean (MPa)','Strength SD (MPa)','Modulus mean (GPa)','Modulus SD (GPa)','Representative trial']
    for j,h in enumerate(barheaders,16): bars.cell(1,j,h)
    metrics=['strength','modulus','break_strain','toughness']
    representatives=[]
    for sample_idx,(name,trials) in enumerate(samples.items(),2):
        eligible=included_trials(trials);stats=sample_statistics(trials)
        rep=None
        if eligible:
            selected=[pair for pair in eligible if pair[1].get('representative',False)]
            rep=selected[0] if selected else min(eligible,key=lambda pair:abs(pair[1]['strength']-stats['strength'][0]))
            representatives.append((name,*rep))
        for j,v in enumerate([name,stats['strength'][0],stats['strength'][1],stats['modulus'][0],stats['modulus'][1],rep[0] if rep else None],16): bars.cell(sample_idx,j,v)
        for label,r in trials:
            summary.append([name,label,r['break_status'],*[r[k] for k in metrics],r['r2'],r['width'],r['thickness'],r['length'],r['fit_min'],r['fit_max'],'Yes' if r.get('included',True) else 'No'])
            ws=wb.create_sheet(f'Trial_{len(wb.worksheets)-3:03d}')
            ws.cell(1,16,name); ws.cell(2,16,label)
            cols=['Time (s)','Force (N)','Stroke (mm)','Strain (%)','Stress (MPa)','Plot stress (MPa)','Modulus fit (MPa)','Fracture marker']
            for j,v in enumerate(cols,16): ws.cell(4,j,v)
            plot,end=plot_values(r,smooth)
            for i in range(len(r['stress'])):
                strain=r['strain'][i]*100
                fit=r['modulus']*1000*r['strain'][i]+r['intercept'] if r['fit_min']<=strain<=r['fit_max'] else None
                vals=[*r['raw'][i],strain,r['stress'][i],plot[i],fit,'Fracture' if i==r['break_index'] else None]
                for j,v in enumerate(vals,16): ws.cell(i+5,j,float(v) if isinstance(v,np.generic) else v)
            chart=ScatterChart(); format_chart(chart,'Strain (%)','Stress (MPa)')
            for ycol,title,count in [(21,'Stress',end),(22,'Modulus fit',len(r['stress']))]:
                series=Series(Reference(ws,min_col=ycol,min_row=5,max_row=4+count),Reference(ws,min_col=19,min_row=5,max_row=4+count),title=title)
                series.graphicalProperties.line.width=12700; chart.series.append(series)
            ws.add_chart(chart,'A1')
        summary.append([name,'Mean','',*[stats[k][0] for k in metrics]])
        summary.append([name,'SD (sample)','',*[stats[k][1] for k in metrics]])
    chart=ScatterChart(); format_chart(chart,'Strain (%)','Stress (MPa)')
    for n,(name,label,r) in enumerate(representatives):
        col=16+n*2; overview.cell(1,col,name); overview.cell(2,col,'Strain (%)'); overview.cell(2,col+1,'Stress (MPa)')
        plot,end=plot_values(r,smooth)
        for i in range(end): overview.cell(i+3,col,float(r['strain'][i]*100)); overview.cell(i+3,col+1,float(plot[i]))
        series=Series(Reference(overview,min_col=col+1,min_row=3,max_row=end+2),Reference(overview,min_col=col,min_row=3,max_row=end+2),title=name)
        series.graphicalProperties.line.width=12700; chart.series.append(series)
    overview.add_chart(chart,'A1')
    main=BarChart(); second=BarChart(); main.grouping='clustered'; second.grouping='clustered'
    count=len(samples)+1
    # Matching empty companion series keep the two value-axis bar groups aligned.
    for j,v in enumerate(['Sample','Strength (MPa)','Spacer 1','Spacer 2','Modulus (GPa)','Strength SD','Modulus SD'],23): bars.cell(1,j,v)
    for i in range(2,count+1):
        values=[bars.cell(i,16).value,bars.cell(i,17).value,None,None,bars.cell(i,19).value,bars.cell(i,18).value,bars.cell(i,20).value]
        for j,v in enumerate(values,23): bars.cell(i,j,v)
    last=count
    main.add_data(Reference(bars,min_col=24,max_col=25,min_row=1,max_row=last),titles_from_data=True)
    second.add_data(Reference(bars,min_col=26,max_col=27,min_row=1,max_row=last),titles_from_data=True)
    for c,series_index,errcol,color in [(main,0,'AB','4472C4'),(second,1,'AC','ED7D31')]:
        c.set_categories(Reference(bars,min_col=23,min_row=2,max_row=last)); c.gapWidth=80;c.overlap=0
        c.series[series_index].graphicalProperties.solidFill=color
        c.series[series_index].graphicalProperties.line.solidFill=color
        c.series[1-series_index].graphicalProperties.noFill=True
        c.series[1-series_index].graphicalProperties.line.noFill=True
        ref=f"'Properties'!${errcol}$2:${errcol}${last}"
        c.series[series_index].errBars=ErrorBars(errDir='y',errBarType='both',errValType='cust',plus=NumDataSource(numRef=NumRef(f=ref)),minus=NumDataSource(numRef=NumRef(f=ref)))
    format_chart(main,'Sample','Tensile strength (MPa)'); format_chart(second,'Sample',"Young's modulus (GPa)")
    second.y_axis.axId=200; second.y_axis.axPos='r'; second.y_axis.crosses='max';second.y_axis.crossAx=10
    main.y_axis.crossAx=10;main.x_axis.crossAx=100;main.x_axis.tickLblPos='low'
    # Reserve both margins for the two explicit value axes in the square figure.
    main.layout=Layout(manualLayout=ManualLayout(layoutTarget='inner',xMode='edge',yMode='edge',wMode='factor',hMode='factor',x=.10,y=.04,w=.80,h=.80))
    main.legend.legendEntry=[LegendEntry(idx=1,delete=True),LegendEntry(idx=2,delete=True)]
    main+=second; bars.add_chart(main,'A1')
    for ws in wb:
        ws.freeze_panes='A2' if ws==summary else 'P5'
        for row in ws:
            for cell in row:
                cell.font=Font(name='Arial',size=14)
                if isinstance(cell.value,(int,float)): cell.number_format='0.0000'
        for cell in ws[1]:
            if cell.value is not None: cell.fill=PatternFill('solid',fgColor='DEEAF6')
        for col in range(1,ws.max_column+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width=24 if col>=16 or ws==summary else 8
    wb.save(path)

def process_files(configs,path,fit_min=.05,fit_max=.25,drop=.5,smooth=False,emit=print):
    samples={}
    for config in configs:
        file,width,thickness,length,overrides=config[:5]
        sample_fit_min,sample_fit_max=config[5:7] if len(config)>=7 else (fit_min,fit_max)
        trial_settings=config[7] if len(config)>7 else {}
        representative=config[8] if len(config)>8 else None
        name=Path(file).stem
        if name in samples: raise ValueError(f'Duplicate sample name: {name}')
        trials=[]
        for label,raw in read_trials(file):
            emit(f'{name}: {label}')
            options=trial_settings.get(label,{})
            result=analyze_trial(raw,width,options.get('thickness',thickness),length,options.get('fit_min',sample_fit_min),options.get('fit_max',sample_fit_max),drop,overrides.get(label))
            result['included']=options.get('included',True)
            result['representative']=label==representative
            trials.append((label,result))
        samples[name]=trials
    export_excel(samples,path,smooth)
    return samples
